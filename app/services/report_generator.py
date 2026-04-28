from __future__ import annotations

import json
from pathlib import Path

try:
    import pandas as pd
except Exception:  # pragma: no cover
    pd = None

try:
    from docx import Document
except Exception:  # pragma: no cover
    Document = None


class ReportGenerator:
    def write_json(self, data: dict, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
        return path

    def write_metrics_xlsx(self, result: dict, path: Path) -> Path | None:
        if pd is None:
            return None
        path.parent.mkdir(parents=True, exist_ok=True)
        metrics = result.get("metrics", {})
        rows = [{"metric": key, "value": value} for key, value in metrics.items()]
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            pd.DataFrame(rows).to_excel(writer, sheet_name="Metrics", index=False)
            pd.DataFrame(result.get("policy_result", {}).get("documents", [])).to_excel(writer, sheet_name="Policy Documents", index=False)
            pd.DataFrame(result.get("retrieval_result", {}).get("queries", [])).drop(columns=["results"], errors="ignore").to_excel(writer, sheet_name="Retrieval Queries", index=False)
        return path

    def write_docx_report(self, result: dict, path: Path) -> Path | None:
        if Document is None:
            return None
        path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_heading("AI4BIZ RAG Benchmark & Compliance Report", level=0)
        doc.add_paragraph(f"Run ID: {result.get('run_id')}")
        doc.add_paragraph(f"Package ID: {result.get('package_id')}")
        doc.add_paragraph(f"Status: {result.get('status')}")
        doc.add_paragraph(f"Final Score: {result.get('final_score')}")
        doc.add_heading("Metric Overview", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        for key, value in result.get("metrics", {}).items():
            row = table.add_row().cells
            row[0].text = str(key)
            row[1].text = str(value)
        doc.add_heading("Executive Recommendation", level=1)
        doc.add_paragraph(result.get("recommendation", "Review weak documents and apply controlled remediation."))
        doc.save(path)
        return path
